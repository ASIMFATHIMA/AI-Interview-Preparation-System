"""
Resume Parsing Module — Stage 1: raw text extraction.

Given a path to an uploaded resume file (.pdf, .docx, or .txt), extract plain text.
Stage 2 (structured extraction via LLM) lives in llm_extractor.py.
"""
import os

import pdfplumber
from docx import Document


class UnsupportedFileTypeError(Exception):
    pass


def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        return _extract_pdf_text(filepath)
    elif ext == ".docx":
        return _extract_docx_text(filepath)
    elif ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    else:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported: .pdf, .docx, .txt"
        )


def _extract_pdf_text(filepath: str) -> str:
    text_parts = []
    with pdfplumber.open(filepath) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_parts.append(page_text)
    return "\n".join(text_parts).strip()


def _extract_docx_text(filepath: str) -> str:
    doc = Document(filepath)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of any tables (common in resume templates)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()
